"""Format-native helper checks use real editable files and durable manifests."""

from __future__ import annotations

import importlib
import json
from pathlib import Path
from types import ModuleType
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from clio_agent.runtime.document_runtime import STACK_ROOT


@pytest.fixture
def helper(monkeypatch: pytest.MonkeyPatch) -> ModuleType:
    """Import the standalone helper just as its interpreter resolves siblings."""
    monkeypatch.syspath_prepend(str(STACK_ROOT))
    return importlib.import_module("cli")


def _content(result: dict[str, Any]) -> dict[str, Any]:
    assert result["status"] == "complete", result
    return json.loads(Path(result["content"]).read_text(encoding="utf-8"))


def test_word_inspection_keeps_body_order_and_original(helper: ModuleType, tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "report résumé.docx"
    document = Document()
    document.add_heading("Reliable report", level=1)
    document.add_table(rows=1, cols=2).cell(0, 1).text = "table value"
    document.add_paragraph("After the table")
    document.save(source)
    original = source.read_bytes()
    result = helper.prepare(source, tmp_path / "output")
    content = _content(result)
    assert [block["type"] for block in content["blocks"]] == ["paragraph", "table", "paragraph"]
    assert content["blocks"][1]["rows"][0][1] == "table value"
    assert source.read_bytes() == original
    assert json.loads(Path(result["manifest"]).read_text(encoding="utf-8")) == result


def test_presentation_inspection_includes_notes_and_slide_selection(
    helper: ModuleType,
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    source = tmp_path / "slides.pptx"
    deck = Presentation()
    for index in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Slide {index + 1}"
        slide.notes_slide.notes_text_frame.text = f"Speaker {index + 1}"
    deck.save(source)
    result = helper.prepare(source, tmp_path / "out", pages="2")
    content = _content(result)
    assert content["slide_count"] == 3
    assert content["selected_slides"] == [2]
    assert content["slides"][0]["notes"] == "Speaker 2"
    assert any(row.get("text") == "Slide 2" for row in content["slides"][0]["shapes"])


def test_workbook_separates_formulas_from_uncalculated_values(
    helper: ModuleType,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook

    source = tmp_path / "budget.xlsx"
    book = Workbook()
    tab = book.active
    assert tab is not None
    tab.title = "Budget Inputs"
    tab.append([2, 3, "=SUM(A1:B1)"])
    book.save(source)
    result = helper.prepare(source, tmp_path / "out", sheet="Budget Inputs", cell_range="A1:C1")
    cells = _content(result)["sheets"][0]["cells"]
    assert cells[2]["formula"] == "=SUM(A1:B1)"
    assert cells[2]["value"] is None
    evidence = helper.workbook_calculation_evidence(source)
    assert evidence["formula_count"] == evidence["issue_count"] == 1
    assert evidence["issues"][0]["error"] == "missing_formula_cache"


def test_large_workbook_refuses_full_read_but_allows_range(
    helper: ModuleType, tmp_path: Path
) -> None:
    from openpyxl import Workbook

    source = tmp_path / "large.xlsx"
    book = Workbook()
    tab = book.active
    assert tab is not None
    tab["A1"] = "first"
    tab["Z1000"] = "last"
    book.save(source)
    result = helper.prepare(source, tmp_path / "out")
    assert result["status"] == "failed"
    assert "select a sheet and cell range" in result["error"]
    selected = helper.prepare(source, tmp_path / "out", sheet="Sheet", cell_range="A1:B2")
    assert _content(selected)["sheets"][0]["cells"][0]["value"] == "first"
    assert selected["run_directory"] != result["run_directory"]


def test_csv_inspection_respects_quotes(helper: ModuleType, tmp_path: Path) -> None:
    source = tmp_path / "quoted.csv"
    source.write_text('name,value\n"two, words",4\n', encoding="utf-8")
    content = _content(helper.prepare(source, tmp_path / "out"))
    assert content["rows"][1] == ["two, words", "4"]


@pytest.mark.parametrize("selection", ["0", "3-2", "1-1000000000", "1-2-3", "x", "6"])
def test_invalid_page_ranges_are_bounded(helper: ModuleType, selection: str) -> None:
    with pytest.raises(RuntimeError, match="Invalid page range"):
        helper.select_pages(selection, 5)


def test_long_document_requires_explicit_range(helper: ModuleType) -> None:
    with pytest.raises(RuntimeError, match="select a range"):
        helper.select_pages("", 200)
    assert helper.select_pages("2-4,3,10", 200) == [2, 3, 4, 10]


def test_recalculation_refuses_external_links_before_converter(
    helper: ModuleType,
    tmp_path: Path,
) -> None:
    source = tmp_path / "linked.xlsx"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    result = helper.prepare(source, tmp_path / "out", action="recalculate")
    assert result["status"] == "failed"
    assert "external links" in result["error"]


def test_cached_empty_string_is_distinct_from_missing_cache(
    helper: ModuleType, tmp_path: Path
) -> None:
    source = tmp_path / "cache.xlsx"
    with ZipFile(source, "w") as archive:
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="str"><f>IF(1=1,"","x")</f><v/></c><c r="B1" t="e"><f>1/0</f><v>#DIV/0!</v></c></row></sheetData></worksheet>',
        )
    evidence = helper.workbook_calculation_evidence(source)
    assert evidence["formula_count"] == 2
    assert evidence["issue_count"] == 1
    assert evidence["issues"][0]["cell"] == "B1"


def test_converter_failure_is_preserved_in_manifest(
    helper: ModuleType,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "broken.xlsx"
    with ZipFile(source, "w"):
        pass

    def fail(source: Path, output: Path, target: str) -> Path:
        raise RuntimeError("converter completed without an output")

    monkeypatch.setattr(helper, "office_convert", fail)
    result = helper.prepare(source, tmp_path / "out", action="recalculate")
    assert result["status"] == "failed"
    assert result["error"] == "converter completed without an output"
    assert json.loads(Path(result["manifest"]).read_text())["status"] == "failed"
