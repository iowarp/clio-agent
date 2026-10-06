"""Exercise creation, editing, extraction, rendering and formula verification on real files."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from cli import prepare
from process import DocumentError


def smoke(root: Path, *, require_office: bool = True) -> dict[str, Any]:
    """Build synthetic fixtures and verify the actual document stack end to end."""
    import pymupdf
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from PIL import Image, ImageDraw
    from pptx import Presentation

    root.mkdir(parents=True, exist_ok=True)
    sources = root / "inputs with spaces"
    sources.mkdir(exist_ok=True)
    image = sources / "synthetic diagram.png"
    canvas = Image.new("RGB", (480, 200), "white")
    drawing = ImageDraw.Draw(canvas)
    drawing.rectangle((30, 40, 170, 160), outline="navy", width=5)
    drawing.line((170, 100, 310, 100), fill="red", width=6)
    drawing.ellipse((310, 40, 450, 160), outline="green", width=5)
    canvas.save(image)
    pdf = sources / "synthetic résumé.pdf"
    with pymupdf.open() as document:
        for index in range(3):
            page = document.new_page()
            if index == 1:
                page.insert_image(pymupdf.Rect(72, 72, 552, 272), filename=str(image))
            else:
                page.insert_text((72, 72), f"Synthetic evidence page {index + 1}")
        document.save(pdf)
    word = sources / "synthetic report.docx"
    document = Document()
    document.add_heading("Synthetic runtime qualification", level=1)
    document.add_paragraph("Before edit")
    document.add_table(rows=2, cols=2).cell(1, 1).text = "Table evidence"
    document.add_picture(str(image))
    document.save(str(word))
    edited = Document(str(word))
    edited.paragraphs[1].runs[0].text = "After edit"
    edited.save(str(word))
    slides = sources / "synthetic slides.pptx"
    deck = Presentation()
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Evidence slide {index + 1}"
        slide.notes_slide.notes_text_frame.text = f"Speaker evidence {index + 1}"
        if index == 1:
            from pptx.util import Inches

            slide.shapes.add_picture(str(image), Inches(1), Inches(2), width=Inches(6))
    deck.save(str(slides))
    workbook = sources / "synthetic budget.xlsx"
    book = Workbook()
    tab = book.active
    if tab is None:
        raise DocumentError("New workbook has no active sheet")
    tab.title = "Budget Inputs"
    tab.append(["First", "Second", "Total"])
    tab.append([2, 3, "=SUM(A2:B2)"])
    book.save(workbook)
    edited_book = load_workbook(workbook)
    edited_book["Budget Inputs"]["A2"] = 4
    edited_book.save(workbook)
    results = []
    for source in (pdf, word, slides, workbook):
        inspected = prepare(source, root / "outputs", action="inspect")
        if inspected["status"] != "complete":
            raise DocumentError(str(inspected))
        if source == pdf:
            extracted = json.loads(Path(inspected["content"]).read_text(encoding="utf-8"))
            if extracted["pages"][1]["text"] or extracted["pages"][1]["image_count"] != 1:
                raise DocumentError("Image-only PDF page was not distinguished from embedded text")
        results.append(inspected)
        if source == pdf or require_office:
            rendered = prepare(source, root / "outputs", action="render")
            if rendered["status"] != "complete" or not rendered["images"]:
                raise DocumentError(str(rendered))
            results.append(rendered)
    if require_office:
        recalculated = prepare(workbook, root / "outputs", action="recalculate")
        if recalculated["status"] != "complete" or recalculated["formula_count"] != 1:
            raise DocumentError(str(recalculated))
        formulas = load_workbook(recalculated["path"], data_only=False)
        values = load_workbook(recalculated["path"], data_only=True)
        if (
            formulas["Budget Inputs"]["C2"].data_type != "f"
            or values["Budget Inputs"]["C2"].value != 7
        ):
            raise DocumentError("Recalculation did not preserve the formula and expected total")
        results.append(recalculated)
    result = {
        "status": "complete",
        "synthetic": True,
        "office_render_verified": require_office,
        "visual_review": "pending",
        "operations": results,
    }
    (root / "smoke.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return result


def main() -> None:
    """Require full Office qualification unless explicitly running the PDF/text subset."""
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("--pdf-text-only", action="store_true")
    args = parser.parse_args()
    result = smoke(args.output.resolve(), require_office=not args.pdf_text_only)
    print(json.dumps({key: value for key, value in result.items() if key != "operations"}))


if __name__ == "__main__":
    main()
